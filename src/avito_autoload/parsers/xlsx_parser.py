"""Parse input xlsx price files into structured rows."""

import logging
from pathlib import Path

from openpyxl import load_workbook

from avito_autoload.models.avito_row import InputRow

logger = logging.getLogger(__name__)

# Mapping: normalized column name -> InputRow field
COLUMN_MAP: dict[str, str] = {
    "код": "code",
    "артикул": "article",
    "номенклатура": "name",
    "наименование": "name",
    "количество": "quantity",
    "стоимость": "cost",
    "цена": "price",
}


def _normalize_header(header: str) -> str:
    """Normalize a column header for matching."""
    return header.strip().lower()


def _detect_columns(header_row: tuple) -> dict[int, str]:
    """Map column indices to InputRow field names."""
    mapping: dict[int, str] = {}
    for idx, cell_value in enumerate(header_row):
        if cell_value is None:
            continue
        normalized = _normalize_header(str(cell_value))
        if normalized in COLUMN_MAP:
            mapping[idx] = COLUMN_MAP[normalized]
    return mapping


def parse_xlsx(file_path: Path, sheet_name: str | None = None) -> list[InputRow]:
    """Read input xlsx and return list of InputRow.

    Expected columns: Код, Артикул, Номенклатура, Количество, Стоимость, Цена.
    The first row is treated as headers. Column names are matched case-insensitively.
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows_iter = ws.iter_rows(values_only=True)

    # First row = headers
    header_row = next(rows_iter, None)
    if header_row is None:
        wb.close()
        return []

    col_map = _detect_columns(header_row)
    if "code" not in col_map.values():
        wb.close()
        raise ValueError(
            f"Column 'Код' not found in headers: {[h for h in header_row if h]}"
        )

    result: list[InputRow] = []
    for row_num, row in enumerate(rows_iter, start=2):
        raw: dict[str, object] = {}
        for idx, field_name in col_map.items():
            value = row[idx] if idx < len(row) else None
            raw[field_name] = value

        # Skip empty rows
        code_val = raw.get("code")
        if code_val is None or str(code_val).strip() == "":
            continue

        # Ensure code is always a string (preserve leading zeros)
        raw["code"] = str(raw["code"]).strip()

        # Defaults for missing fields
        if raw.get("article") is None:
            raw["article"] = ""
        else:
            raw["article"] = str(raw["article"]).strip()

        if raw.get("name") is None:
            raw["name"] = ""
        else:
            raw["name"] = str(raw["name"]).strip()

        raw["quantity"] = int(raw.get("quantity") or 0)
        raw["cost"] = float(raw.get("cost") or 0.0)
        raw["price"] = float(raw.get("price") or 0.0)

        try:
            result.append(InputRow(**raw))
        except Exception:
            logger.warning("Skipping row %d: invalid data %s", row_num, raw)

    wb.close()
    logger.info("Parsed %d rows from %s", len(result), file_path.name)
    return result
