"""Target audience (CA) analysis service via LLM."""

import asyncio
import json
import logging

from avito_autoload.categorizers.llm_categorizer import _call_anthropic, _parse_llm_response

logger = logging.getLogger(__name__)


def _build_ca_prompt(niche: str, competitor_samples: list[str]) -> str:
    """Build prompt for target audience analysis."""
    samples_text = ""
    if competitor_samples:
        samples_text = "\n\nПримеры объявлений конкурентов:\n"
        for i, sample in enumerate(competitor_samples[:20], 1):
            # Truncate very long descriptions
            text = sample[:500] + "..." if len(sample) > 500 else sample
            samples_text += f"{i}. {text}\n"

    return f"""Ты — маркетолог-аналитик. Проведи анализ целевой аудитории для ниши на Авито.

Ниша: «{niche}»
{samples_text}

Проведи детальный анализ и ответь СТРОГО в формате JSON:
{{
  "segments": [
    {{"name": "...", "description": "...", "share_percent": 40}},
    {{"name": "...", "description": "...", "share_percent": 35}},
    {{"name": "...", "description": "...", "share_percent": 25}}
  ],
  "main_pain": "Главная боль целевой аудитории — одним предложением",
  "pain_points": [
    "Боль 1",
    "Боль 2",
    "Боль 3",
    "Боль 4",
    "Боль 5"
  ],
  "top_usp": [
    "УТП 1 для объявлений",
    "УТП 2 для объявлений",
    "УТП 3 для объявлений",
    "УТП 4 для объявлений",
    "УТП 5 для объявлений"
  ],
  "objections": [
    "Возражение 1 и как снять",
    "Возражение 2 и как снять",
    "Возражение 3 и как снять"
  ],
  "keywords": ["ключевое слово 1", "ключевое слово 2", "..."],
  "tone": "Рекомендованный тон коммуникации",
  "full_analysis": "Полный текст анализа в формате Markdown (2000-3000 слов). Включи: описание ниши, портреты сегментов, психографику, путь клиента, мотивации, страхи, возражения, рекомендации по коммуникации."
}}

Ответь ТОЛЬКО валидным JSON."""


def _parse_competitor_descriptions(file_path: str) -> list[str]:
    """Extract description texts from competitor XLSX file."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True)
        ws = wb.active

        descriptions = []
        # Try to find description column
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            wb.close()
            return []

        desc_col = None
        for i, cell in enumerate(header_row):
            if cell and any(
                kw in str(cell).lower()
                for kw in ["описание", "description", "текст", "text"]
            ):
                desc_col = i
                break

        if desc_col is None:
            # Fallback: try title/name column
            for i, cell in enumerate(header_row):
                if cell and any(
                    kw in str(cell).lower()
                    for kw in ["название", "title", "заголовок", "name"]
                ):
                    desc_col = i
                    break

        if desc_col is not None:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if desc_col < len(row) and row[desc_col]:
                    descriptions.append(str(row[desc_col]))

        wb.close()
        return descriptions
    except Exception:
        logger.exception("Failed to parse competitor file")
        return []


async def analyze_target_audience(
    niche: str,
    competitors_path: str | None = None,
) -> dict:
    """Run CA analysis via LLM.

    Returns dict with keys: segments, main_pain, pain_points, top_usp,
    objections, keywords, tone, full_analysis.
    """
    competitor_samples = []
    if competitors_path:
        competitor_samples = await asyncio.to_thread(
            _parse_competitor_descriptions, competitors_path
        )

    prompt = _build_ca_prompt(niche, competitor_samples)
    response = await asyncio.to_thread(_call_anthropic, prompt)

    if response is None:
        return {
            "segments": [],
            "main_pain": "Не удалось провести анализ",
            "pain_points": [],
            "top_usp": [],
            "objections": [],
            "keywords": [],
            "tone": "",
            "full_analysis": "",
        }

    # Parse JSON response
    text = response.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        lines = [ln for ln in lines if not ln.strip().startswith("```")]
        text = "\n".join(lines)

    try:
        result = json.loads(text)
        return result
    except json.JSONDecodeError:
        logger.warning("Failed to parse CA analysis as JSON, returning raw text")
        return {
            "segments": [],
            "main_pain": "",
            "pain_points": [],
            "top_usp": [],
            "objections": [],
            "keywords": [],
            "tone": "",
            "full_analysis": text,
        }


def format_ca_summary(ca_data: dict, niche: str) -> str:
    """Format CA analysis as a Telegram message."""
    segments = ca_data.get("segments", [])
    main_pain = ca_data.get("main_pain", "")
    top_usp = ca_data.get("top_usp", [])

    lines = [
        "\u2705 Этап 1/4: Анализ ЦА завершён",
        "",
        f"\U0001F4CA Ниша: {niche}",
        "",
        "\U0001F465 Основные сегменты покупателей:",
        "",
    ]

    for i, seg in enumerate(segments, 1):
        name = seg.get("name", "")
        desc = seg.get("description", "")
        lines.append(f"{i}. <b>{name}</b> \u2014 {desc}")

    if main_pain:
        lines.extend(["", f"\U0001F3AF Главная боль ЦА:\n{main_pain}"])

    if top_usp:
        lines.extend(["", "\U0001F4A1 Топ-5 УТП для объявлений:"])
        for i, usp in enumerate(top_usp[:5], 1):
            lines.append(f"{i}. {usp}")

    lines.extend(["", "Подробный анализ \u2014 в прикреплённом файле."])

    return "\n".join(lines)
