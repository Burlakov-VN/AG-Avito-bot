"""Competitor analysis service — extracts patterns from competitor listings."""

import asyncio
import json
import logging

from avito_autoload.categorizers.llm_categorizer import _call_anthropic

logger = logging.getLogger(__name__)


def _read_competitor_data(file_path: str) -> list[dict]:
    """Read competitor XLSX and extract structured data."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True)
        ws = wb.active

        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            wb.close()
            return []

        headers = [str(h).strip().lower() if h else "" for h in header_row]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            entry = {}
            for i, val in enumerate(row):
                if i < len(headers) and headers[i] and val is not None:
                    entry[headers[i]] = str(val)
            if entry:
                rows.append(entry)

        wb.close()
        return rows
    except Exception:
        logger.exception("Failed to read competitor file")
        return []


def _build_competitor_prompt(niche: str, competitor_data: list[dict]) -> str:
    """Build prompt for competitor analysis."""
    # Sample up to 30 competitor entries
    samples = competitor_data[:30]
    samples_text = json.dumps(samples, ensure_ascii=False, indent=2)

    return f"""Ты — маркетолог-аналитик. Проведи анализ объявлений конкурентов на Авито.

Ниша: «{niche}»

Данные конкурентов (до 30 объявлений):
{samples_text}

Проанализируй и ответь СТРОГО в формате JSON:
{{
  "total_analyzed": {len(competitor_data)},
  "avg_description_length": 150,
  "patterns": [
    "Паттерн 1 — что часто используют конкуренты",
    "Паттерн 2",
    "Паттерн 3"
  ],
  "best_practices": [
    "Лучшая практика 1 — что стоит перенять",
    "Лучшая практика 2",
    "Лучшая практика 3"
  ],
  "weaknesses": [
    "Слабость 1 — что конкуренты делают плохо",
    "Слабость 2",
    "Слабость 3"
  ],
  "emojis": ["✅", "🔹", "⚡️", "📌"],
  "optimal_length_words": [150, 300],
  "key_phrases": ["фраза 1", "фраза 2", "фраза 3"],
  "pricing_insights": "Наблюдения по ценообразованию"
}}

Ответь ТОЛЬКО валидным JSON."""


async def analyze_competitors(
    niche: str,
    competitors_path: str,
) -> dict:
    """Analyze competitor listings via LLM.

    Returns dict with patterns, best_practices, weaknesses, emojis, etc.
    """
    competitor_data = await asyncio.to_thread(_read_competitor_data, competitors_path)

    if not competitor_data:
        return {
            "total_analyzed": 0,
            "patterns": [],
            "best_practices": [],
            "weaknesses": [],
            "emojis": ["\u2705", "\ud83d\udd39", "\u26a1\ufe0f", "\ud83d\udccc"],
            "optimal_length_words": [150, 300],
            "key_phrases": [],
        }

    prompt = _build_competitor_prompt(niche, competitor_data)
    response = await asyncio.to_thread(_call_anthropic, prompt)

    if response is None:
        return {
            "total_analyzed": len(competitor_data),
            "patterns": [],
            "best_practices": [],
            "weaknesses": [],
            "emojis": ["\u2705", "\ud83d\udd39", "\u26a1\ufe0f", "\ud83d\udccc"],
            "optimal_length_words": [150, 300],
            "key_phrases": [],
        }

    text = response.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        lines = [ln for ln in lines if not ln.strip().startswith("```")]
        text = "\n".join(lines)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        logger.warning("Failed to parse competitor analysis JSON")
        return {
            "total_analyzed": len(competitor_data),
            "patterns": [],
            "best_practices": [],
            "weaknesses": [],
            "emojis": ["\u2705", "\ud83d\udd39", "\u26a1\ufe0f", "\ud83d\udccc"],
            "optimal_length_words": [150, 300],
            "key_phrases": [],
        }
